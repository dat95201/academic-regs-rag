"""Turn files in data/raw into Markdown in data/processed, the input of chunking.

- Scanned PDF (no text layer) -> OCR (app.ingest.ocr). The result should be reviewed by hand.
- PDF with a text layer, DOCX, MD, TXT -> text extraction with page markers.
- An existing data/processed/<doc_id>.md is never overwritten: it may hold manual fixes.

TODO(Task 2): drop repeated headers/footers and detect legacy TCVN3/VNI fonts for text PDFs.
"""

from __future__ import annotations

import hashlib
import json
import logging
import re
import unicodedata
from pathlib import Path

log = logging.getLogger("parse")

SUPPORTED = {".pdf", ".docx", ".md", ".txt"}
# Characters typical of text typed in TCVN3/VNI fonts and extracted without conversion.
LEGACY_FONT_HINT = re.compile(r"[µ¸¶·¹¨»¾¼½Æ©ÇÊÈÉË®ÌÐÎÏÑªÒÕÓÔÖ×ÝØÜÞßãáâäå«æçèéêëìíîïóðñòôõøöùúûüþ]{3,}")


def doc_id_for(path: Path) -> str:
    s = unicodedata.normalize("NFD", path.stem).replace("đ", "d").replace("Đ", "D")
    s = "".join(c for c in s if unicodedata.category(c) != "Mn").lower()
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def header(path: Path, **extra) -> str:
    meta = {"doc_id": doc_id_for(path), "source_file": unicodedata.normalize("NFC", path.name),
            "sha256": sha256(path), **extra}
    return f"<!-- doc: {json.dumps(meta, ensure_ascii=False)} -->\n"


def clean(text: str) -> str:
    text = unicodedata.normalize("NFC", text)
    text = re.sub(r"[ \t]+", " ", text)
    return "\n".join(line.strip() for line in text.splitlines())


def pdf_to_markdown(path: Path) -> str:
    import pymupdf

    from app.ingest import ocr

    doc = pymupdf.open(path)
    if not ocr.has_text_layer(doc):
        log.warning("%s: scanned PDF, running OCR (%d pages); review the result before relying on it",
                    path.name, doc.page_count)
        return ocr.ocr_pdf(path, range(doc.page_count))
    out = [header(path, extractor="pymupdf")]
    for i, page in enumerate(doc):
        text = clean(page.get_text())
        lines = text.splitlines()
        if lines and re.fullmatch(r"\d{1,3}", lines[0].strip()):
            lines = lines[1:]  # page number
        out.append(f"\n<!-- page: {i + 1} -->\n" + "\n".join(lines))
    return "\n".join(out)


def docx_to_markdown(path: Path) -> str:
    import docx  # python-docx

    document = docx.Document(str(path))
    body = [clean(p.text) for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        rows = [[clean(c.text).replace("\n", " ") for c in row.cells] for row in table.rows]
        if rows:
            body.append("")
            body.append("| " + " | ".join(rows[0]) + " |")
            body.append("|" + "---|" * len(rows[0]))
            body.extend("| " + " | ".join(r) + " |" for r in rows[1:])
            body.append("")
    return header(path, extractor="python-docx") + "\n<!-- page: 1 -->\n" + "\n".join(body)


def text_to_markdown(path: Path) -> str:
    return header(path, extractor="text") + "\n<!-- page: 1 -->\n" + clean(path.read_text(encoding="utf-8"))


def prepare(raw_dir: Path, processed_dir: Path) -> list[Path]:
    """Make sure every raw file has a processed Markdown; return the Markdown files to ingest."""
    processed_dir.mkdir(parents=True, exist_ok=True)
    for raw in sorted(raw_dir.iterdir()) if raw_dir.exists() else []:
        if raw.name.startswith(".") or raw.is_dir():
            continue
        if raw.suffix.lower() not in SUPPORTED:
            log.warning("%s: unsupported file type, skipped", raw.name)
            continue
        target = processed_dir / f"{doc_id_for(raw)}.md"
        if target.exists():
            continue
        convert = {".pdf": pdf_to_markdown, ".docx": docx_to_markdown}.get(
            raw.suffix.lower(), text_to_markdown
        )
        markdown = convert(raw)
        if LEGACY_FONT_HINT.search(markdown):
            log.warning("%s: text looks like a TCVN3/VNI font; convert it to Unicode first", raw.name)
        target.write_text(markdown, encoding="utf-8")
        log.info("%s -> %s", raw.name, target)
    return sorted(p for p in processed_dir.glob("*.md") if ".pages-" not in p.name)
